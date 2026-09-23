"""Render assessment findings into report files in multiple formats.

Supported formats: ``json``, ``csv``, ``md``, ``html``, ``xlsx``, ``pdf``.
CSV/JSON/MD/HTML use the standard library; XLSX uses ``openpyxl`` and PDF uses
``fpdf2`` (both pure-Python, bundle-friendly).
"""
from __future__ import annotations

import csv
import html
import json
import time
from pathlib import Path
from typing import Any

from .exceptions import EngineError

FORMATS = ("json", "csv", "md", "html", "xlsx", "pdf")

_COLUMNS = ("severity", "title", "source", "target", "description", "created")
_SEV_HEX = {
    "high": "C0392B",
    "warning": "C07A00",
    "info": "0A64AD",
    "secure": "1E7E34",
    "hotspot": "C07A00",
    "note": "555555",
}
_SEV_ORDER = {"high": 0, "warning": 1, "hotspot": 2, "info": 3, "note": 4, "secure": 5}


def generate(
    findings: list[dict[str, Any]],
    fmt: str,
    out_path: Path,
    *,
    meta: dict[str, Any] | None = None,
) -> Path:
    """Render ``findings`` to ``out_path`` in ``fmt`` and return the path."""
    fmt = fmt.lower()
    if fmt not in FORMATS:
        raise EngineError("report", f"Unsupported format {fmt!r}. Use one of {FORMATS}.")
    meta = meta or {}
    meta.setdefault("generated", time.strftime("%Y-%m-%d %H:%M:%S"))
    findings = sorted(findings, key=lambda f: _SEV_ORDER.get(f.get("severity", "info"), 9))
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    {
        "json": _json,
        "csv": _csv,
        "md": _md,
        "html": _html,
        "xlsx": _xlsx,
        "pdf": _pdf,
    }[fmt](findings, meta, out_path)
    return out_path


def _counts(findings: list[dict]) -> dict[str, int]:
    c: dict[str, int] = {}
    for f in findings:
        c[f.get("severity", "info")] = c.get(f.get("severity", "info"), 0) + 1
    return c


def _json(findings, meta, out: Path) -> None:
    out.write_text(json.dumps({"meta": meta, "findings": findings}, indent=2, default=str))


def _csv(findings, meta, out: Path) -> None:
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(_COLUMNS)
        for f in findings:
            w.writerow([str(f.get(c, "")).replace("\n", " ") for c in _COLUMNS])


def _md(findings, meta, out: Path) -> None:
    lines = [
        "# Security Assessment Report",
        "",
        f"Generated: {meta.get('generated')}  ·  Tool: msrf {meta.get('tool_version', '')}",
        "",
        "## Summary",
        "",
    ]
    counts = _counts(findings)
    for sev in ("high", "warning", "info", "secure", "hotspot", "note"):
        if counts.get(sev):
            lines.append(f"- **{sev.upper()}**: {counts[sev]}")
    lines += ["", "## Findings", "", "| Severity | Title | Source | Target |", "|---|---|---|---|"]
    for f in findings:
        lines.append(
            f"| {f.get('severity')} | {_md_cell(f.get('title'))} | "
            f"{f.get('source')} | {_md_cell(f.get('target'))} |"
        )
    lines += ["", "## Details", ""]
    for f in findings:
        lines.append(f"### [{f.get('severity')}] {f.get('title')}")
        lines.append(f"- Source: {f.get('source')}  ·  Target: {f.get('target')}")
        if f.get("description"):
            lines.append("")
            lines.append(str(f["description"]))
        lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")


def _md_cell(text: Any) -> str:
    return " ".join(str(text or "").split()).replace("|", "\\|")


def _html(findings, meta, out: Path) -> None:
    counts = _counts(findings)
    summary = " ".join(
        f'<span class="pill {s}">{s.upper()}: {counts[s]}</span>'
        for s in ("high", "warning", "info", "secure", "hotspot", "note")
        if counts.get(s)
    )
    rows = "\n".join(
        f"<tr><td class='sev {html.escape(f.get('severity',''))}'>{html.escape(f.get('severity',''))}</td>"
        f"<td>{html.escape(str(f.get('title','')))}</td>"
        f"<td>{html.escape(str(f.get('source','')))}</td>"
        f"<td>{html.escape(str(f.get('target','')))}</td>"
        f"<td>{html.escape(str(f.get('description','')))}</td></tr>"
        for f in findings
    )
    css = "".join(f".sev.{k}{{color:#{v};font-weight:bold}}" for k, v in _SEV_HEX.items())
    pill = "".join(f".pill.{k}{{background:#{v}}}" for k, v in _SEV_HEX.items())
    out.write_text(
        f"""<!doctype html><html><head><meta charset="utf-8"><title>msrf report</title>
<style>
body{{font-family:'Segoe UI',Arial,sans-serif;margin:24px;color:#1a1a1a}}
h1{{color:#0a64ad}} table{{border-collapse:collapse;width:100%}}
td,th{{border:1px solid #c0c0c0;padding:6px 8px;text-align:left;vertical-align:top}}
th{{background:#eef4fa}} .pill{{color:#fff;padding:2px 8px;border-radius:10px;margin-right:6px;font-size:12px}}
{css}{pill}
</style></head><body>
<h1>Security Assessment Report</h1>
<p>Generated: {html.escape(meta.get('generated',''))} &middot; Tool: msrf {html.escape(str(meta.get('tool_version','')))}</p>
<p>{summary}</p>
<table><tr><th>Severity</th><th>Title</th><th>Source</th><th>Target</th><th>Description</th></tr>
{rows}
</table></body></html>""",
        encoding="utf-8",
    )


def _xlsx(findings, meta, out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    headers = [c.capitalize() for c in _COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF4FA")
    for f in findings:
        ws.append([str(f.get(c, "")) for c in _COLUMNS])
        sev = f.get("severity", "info")
        ws.cell(row=ws.max_row, column=1).font = Font(
            bold=True, color=_SEV_HEX.get(sev, "000000")
        )
    for i, width in enumerate((12, 50, 12, 28, 70, 20), start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    summary = wb.create_sheet("Summary")
    summary.append(["Severity", "Count"])
    for sev, n in _counts(findings).items():
        summary.append([sev, n])
    wb.save(out)


def _pdf(findings, meta, out: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Security Assessment Report", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 6, f"Generated: {meta.get('generated')}   Tool: msrf {meta.get('tool_version','')}",
             new_x="LMARGIN", new_y="NEXT")
    counts = _counts(findings)
    summary = "   ".join(f"{s.upper()}: {n}" for s, n in counts.items())
    pdf.cell(0, 6, summary, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    widths = (22, 90, 22, 45, 90)
    heads = ("Severity", "Title", "Source", "Target", "Description")
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(238, 244, 250)
    for w, h in zip(widths, heads, strict=False):
        pdf.cell(w, 7, h, border=1, fill=True)
    pdf.ln()
    pdf.set_font("Helvetica", "", 8)
    for f in findings:
        cells = [
            str(f.get("severity", "")),
            _fit(f.get("title"), 60),
            str(f.get("source", "")),
            _fit(f.get("target"), 28),
            _fit(f.get("description"), 60),
        ]
        for w, text in zip(widths, cells, strict=False):
            pdf.cell(w, 6, text, border=1)
        pdf.ln()
    pdf.output(str(out))


def _fit(text: Any, n: int) -> str:
    s = " ".join(str(text or "").split())
    # Latin-1 for the core PDF fonts.
    s = s.encode("latin-1", "replace").decode("latin-1")
    return s if len(s) <= n else s[: n - 1] + "…".encode("latin-1", "replace").decode("latin-1")
